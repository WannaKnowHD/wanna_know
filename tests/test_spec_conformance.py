"""기능명세서(docs/기능명세서.xlsx)에 적힌 클래스·메소드가 코드에 모두 있는지 확인.

실행: pytest
명세서를 고치면 이 테스트가 코드와의 차이를 알려 줌.
"""
import importlib.util
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
SPEC = ROOT / "docs" / "기능명세서.xlsx"

FILE_MAP = {
    "llm.py": "chatbot/rag/llm.py",
    "config.py": "chatbot/core/config.py",
    "auth.py": "chatbot/core/auth.py",
    "gradio_app.py": "chatbot/ui/gradio_app.py",
    "logger.py": "chatbot/core/logger.py",
    "db.py": "chatbot/db/db.py",
    "schemas.py": "chatbot/core/schemas.py",
    "document_processor.py": "chatbot/rag/document_processor.py",
    "data/": "chatbot/rag/glossary.py",
    "glossary.py": "chatbot/rag/glossary.py",
    "build_index.py": "build_index.py",
    "retriever.py": "chatbot/rag/retriever.py",
    "vision.py": "chatbot/vision/vision.py",
    "pipeline.py": "chatbot/pipeline.py",
    "validator.py": "chatbot/rag/validator.py",
    "safety.py": "chatbot/core/safety.py",
    "state.py": "chatbot/core/state.py",
    "eval/run_eval.py": "eval/run_eval.py",
    "dashboard.py": "chatbot/ui/dashboard.py",
}


def spec_rows():
    if not SPEC.exists():  # docs/는 커밋하지 않으므로, 명세서를 받아 두지 않았으면 건너뜀
        return []
    wb = openpyxl.load_workbook(SPEC)
    rows = []
    for ws in wb.worksheets:
        if "메소드 명세서" not in ws.title:
            continue
        f = c = None
        for r in range(3, ws.max_row + 1):
            f = (ws[f"B{r}"].value or f)
            c = (ws[f"C{r}"].value or c)
            d = ws[f"D{r}"].value
            if d:
                rows.append((ws.title, f.strip(), c.strip(), d.strip()))
    return rows


def load(path):
    name = "spec_" + path.replace("/", "_").removesuffix(".py")
    if name in sys.modules:
        return sys.modules[name]
    spec = importlib.util.spec_from_file_location(name, ROOT / path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module  # dataclass 등이 모듈을 찾을 수 있게 먼저 등록
    spec.loader.exec_module(module)
    return module


@pytest.mark.skipif(not SPEC.exists(), reason="docs/기능명세서.xlsx가 없음 (드라이브에서 받아 docs/에 넣으면 검사)")
@pytest.mark.parametrize("sheet,file,cls,method", spec_rows() or [("-", "-", "-", "-")])
def test_spec_item_exists(sheet, file, cls, method):
    assert file in FILE_MAP, f"[{sheet}] 명세서 파일명 {file!r}의 경로가 FILE_MAP에 없음"
    module = load(FILE_MAP[file])
    assert hasattr(module, cls), f"{FILE_MAP[file]}에 클래스 {cls} 없음"
    if method.startswith("("):  # (데이터 클래스), (pydantic 모델)
        return
    assert hasattr(getattr(module, cls), method), f"{cls}에 메소드 {method} 없음"
